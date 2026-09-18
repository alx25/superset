import math
import re
import secrets
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field
from superset_core.mcp.decorators import tool

from . import report_cache
from .compare_periods import _build_comparison, _is_additive_metric
from .query_dataset import (
    DatasetFilter,
    JinjaFilterOverride,
    _apply_post_filter,
    _jinja_filters_context,
    _parse_filter,
    _parse_metric,
    _parse_orderby,
    _sanitize_rows,
)
from .sql_analysis import ExtraTable, execute_sql_analysis

_EXPORT_DIR = Path("/home/imercados/.superset/mcp_exports")
_EXPORT_MAX_AGE_SEC = 7200  # 2 horas


def _cleanup_old_exports() -> None:
    if not _EXPORT_DIR.exists():
        return
    cutoff = time.time() - _EXPORT_MAX_AGE_SEC
    for f in _EXPORT_DIR.glob("*.xlsx"):
        try:
            if f.stat().st_mtime < cutoff:
                f.unlink()
        except OSError:
            pass


def _get_export_base_url() -> str:
    try:
        from flask import current_app
        base = current_app.config.get("MCP_EXPORT_BASE_URL", "").rstrip("/")
        if base:
            return base
    except RuntimeError:
        pass
    return "http://192.168.76.11"


def _safe_sheet_name(name: str, used: set[str]) -> str:
    """Excel prohíbe : \\ / ? * [ ] en nombres de hoja y limita a 31 caracteres.
    Agrega un sufijo numérico si el nombre ya está en uso."""
    cleaned = re.sub(r'[:\\/?*\[\]]', "-", str(name)).strip() or "Hoja"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base[: 31 - len(str(suffix)) - 1]}~{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


_METRIC_NAME_RE = re.compile(
    r'(?:SUM|COUNT|AVG|MIN|MAX)\s*\(\s*"?(\w+)"?\s*\)\s*$', re.I
)


def _metric_display_name(metric_label: str) -> str:
    """Nombre legible de una métrica cruda tipo 'AVG(precio)' -> 'precio' —
    para usar en encabezados de columna del Excel de comparación en vez del
    nombre técnico completo. Si no matchea el patrón de agregación simple
    (ej. ya es un alias, o una métrica guardada), se devuelve tal cual."""
    m = _METRIC_NAME_RE.match(metric_label.strip())
    return m.group(1) if m else metric_label


def _format_comparison_rows(
    combined: list[dict[str, Any]],
    metric_labels: list[str],
    period_a_label: str,
    period_b_label: str,
) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    """Renombra las columnas crudas de _build_comparison (ej. 'AVG(precio)_a',
    'AVG(precio)_delta') a encabezados descriptivos usando las etiquetas de
    período reales — genérico para cualquier métrica/etiqueta, sin hardcodear
    nombres de columna. Devuelve (rows, var_pct_cols, delta_cols) — ambas
    listas de encabezados ya renombrados, para que _write_data_sheet les
    aplique el formato numérico y condicional correspondiente."""
    var_pct_cols: list[str] = []
    delta_cols: list[str] = []
    col_map: dict[str, str] = {}
    for label in metric_labels:
        display = _metric_display_name(label)
        col_map[f"{label}_a"] = f"{display} · {period_a_label}"
        col_map[f"{label}_b"] = f"{display} · {period_b_label}"
        var_header = f"{display} · variación: {period_b_label} vs {period_a_label}"
        delta_header = f"{display} · cambio: {period_b_label} − {period_a_label}"
        col_map[f"{label}_var_pct"] = var_header
        col_map[f"{label}_delta"] = delta_header
        col_map[f"{label}_contribution_pct"] = f"{display} · contribución %"
        var_pct_cols.append(var_header)
        delta_cols.append(delta_header)

    rows = [{col_map.get(k, k): v for k, v in row.items()} for row in combined]
    return rows, var_pct_cols, delta_cols


def _write_data_sheet(
    ws: Any,
    df: Any,
    var_pct_cols: list[str],
    delta_cols: list[str] | None = None,
    highlight_col: str | None = None,
) -> None:
    """Escribe encabezados + filas + formato en una hoja YA CREADA.

    'delta_cols' recibe el mismo color condicional (verde/rojo/neutro) que
    'var_pct_cols', pero con formato numérico normal en vez de porcentaje.
    'highlight_col', si se pasa, resalta con un borde la fila de MAYOR y
    MENOR valor de esa columna (ej. mayor aumento / mayor disminución) —
    se recalcula sobre el DataFrame de ESTA hoja, así cada pestaña de un
    split_by resalta su propio extremo local en vez de uno global que ya
    no correspondería a las filas presentes en esa hoja."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    df = df.reset_index(drop=True)
    ws.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="366092")
    header_font = Font(bold=True, color="FFFFFF")
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    highlight_border = Border(
        top=Side(style="medium", color="366092"),
        bottom=Side(style="medium", color="366092"),
    )
    delta_cols = delta_cols or []

    highlight_rows: set[int] = set()
    if highlight_col and highlight_col in df.columns:
        valid = df[highlight_col].dropna()
        if not valid.empty:
            highlight_rows = {int(valid.idxmax()), int(valid.idxmin())}

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(1, col_idx, col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_idx, row_dict in enumerate(df.to_dict("records"), 2):
        is_highlight = (row_idx - 2) in highlight_rows
        for col_idx, col_name in enumerate(df.columns, 1):
            raw = row_dict.get(col_name)
            safe_val = (
                None
                if isinstance(raw, float) and (math.isnan(raw) or math.isinf(raw))
                else raw
            )
            if col_name in var_pct_cols and safe_val is not None:
                try:
                    v = float(safe_val)
                    cell = ws.cell(row_idx, col_idx, v / 100)
                    cell.number_format = "+0.00%;-0.00%;0.00%"
                    cell.fill = green_fill if v > 0 else red_fill if v < 0 else PatternFill()
                except (TypeError, ValueError):
                    cell = ws.cell(row_idx, col_idx, safe_val)
            elif col_name in delta_cols and safe_val is not None:
                try:
                    v = float(safe_val)
                    cell = ws.cell(row_idx, col_idx, v)
                    cell.number_format = "+#,##0.00;-#,##0.00;0.00"
                    cell.fill = green_fill if v > 0 else red_fill if v < 0 else PatternFill()
                except (TypeError, ValueError):
                    cell = ws.cell(row_idx, col_idx, safe_val)
            else:
                cell = ws.cell(row_idx, col_idx, safe_val)
                if isinstance(safe_val, float):
                    cell.number_format = "#,##0.00"
            if is_highlight:
                cell.border = highlight_border

    for col_idx, col_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(col_idx)
        try:
            data_max = int(df[col_name].astype(str).str.len().max() or 10)
        except Exception:
            data_max = 10
        ws.column_dimensions[col_letter].width = min(
            max(len(col_name), data_max) + 3, 55
        )

    ws.auto_filter.ref = ws.dimensions


_BOLD_RUN_RE = re.compile(r"\*\*([^*]+)\*\*")
_SUMMARY_COL_WIDTH = 110
# openpyxl no autoajusta la altura de fila para wrap_text=True (a diferencia
# de Excel al editar a mano) — sin fijarla, un párrafo largo queda con la
# altura default de UNA línea y el resto del texto se corta/superpone con la
# fila siguiente. 0.9 es una aproximación conservadora (de menos) de cuántos
# caracteres entran por línea visible a este ancho de columna con Calibri 11,
# para sobrestimar la altura antes que quedarse corta.
_SUMMARY_CHARS_PER_LINE = int(_SUMMARY_COL_WIDTH * 0.9)
_SUMMARY_LINE_HEIGHT_PT = 15.0


def _summary_rich_text(text: str) -> Any:
    """Convierte '**negrita**' en runs de negrita real (rich text) en vez de
    solo quitar los asteriscos — así el énfasis del LLM se ve en el Excel."""
    if "**" not in text:
        return text
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    parts = _BOLD_RUN_RE.split(text)  # alterna: plano, negrita, plano, negrita...
    blocks: list[Any] = [
        TextBlock(InlineFont(b=True), part) if i % 2 == 1 else part
        for i, part in enumerate(parts)
        if part
    ]
    return CellRichText(*blocks) if blocks else text


def _write_summary_sheet(ws: Any, summary_text: str, title: str) -> None:
    """Escribe una hoja de texto con el análisis — soporta un subconjunto simple
    de markdown: '#'/'##'/'###' como encabezados en negrita, '-'/'*' o '1.'
    como listas, y '**texto**' como negrita real. Calcula la altura de cada
    fila a partir del wrap real del párrafo (ver _SUMMARY_CHARS_PER_LINE)."""
    import textwrap

    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = _SUMMARY_COL_WIDTH
    ws.sheet_view.showGridLines = False

    title_cell = ws.cell(1, 1, title)
    title_cell.font = Font(bold=True, size=14, color="366092")
    ws.row_dimensions[1].height = 22
    row_idx = 3

    def write_paragraph(row: int, text: str, indent: str = "") -> None:
        combined = indent + text
        cell = ws.cell(row, 1, _summary_rich_text(combined))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        plain = _BOLD_RUN_RE.sub(r"\1", combined)
        n_lines = len(
            textwrap.wrap(
                plain, _SUMMARY_CHARS_PER_LINE,
                break_long_words=False, break_on_hyphens=False,
            )
        ) or 1
        ws.row_dimensions[row].height = n_lines * _SUMMARY_LINE_HEIGHT_PT

    heading_sizes = {1: 15, 2: 13, 3: 12}
    prev_was_blank = True
    for raw_line in summary_text.splitlines():
        stripped = raw_line.strip()
        if not stripped:
            if not prev_was_blank:
                row_idx += 1
                prev_was_blank = True
            continue
        prev_was_blank = False

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered_match = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)

        if heading_match:
            if row_idx > 3:
                row_idx += 1  # aire antes de cada encabezado (salvo el primero)
            level = len(heading_match.group(1))
            text = _BOLD_RUN_RE.sub(r"\1", heading_match.group(2))
            cell = ws.cell(row_idx, 1, text)
            cell.font = Font(bold=True, size=heading_sizes.get(level, 12), color="366092")
            ws.row_dimensions[row_idx].height = 20
        elif bullet_match:
            write_paragraph(row_idx, bullet_match.group(1), indent="•  ")
        elif numbered_match:
            write_paragraph(
                row_idx, numbered_match.group(2),
                indent=f"{numbered_match.group(1)}.  ",
            )
        else:
            write_paragraph(row_idx, stripped)
        row_idx += 1


def _add_chart_sheet(
    wb: Any,
    df: Any,
    chart_type: str,
    x_col: str | None,
    y_col: str | None,
    used_names: set[str],
) -> str | None:
    """Crea una hoja 'Gráfico' con un gráfico nativo de Excel (interactivo) a
    partir de las top 20 filas (por valor de y_col). Retorna None si no hay
    columnas utilizables para graficar."""
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    numeric_cols = [c for c in df.columns if str(df[c].dtype).startswith(("float", "int"))]
    non_numeric_cols = [c for c in df.columns if c not in numeric_cols]

    resolved_y = y_col if y_col in numeric_cols else (numeric_cols[0] if numeric_cols else None)
    resolved_x = x_col if x_col in df.columns else (non_numeric_cols[0] if non_numeric_cols else None)
    if resolved_y is None or resolved_x is None:
        return None

    chart_df = df[[resolved_x, resolved_y]].dropna(subset=[resolved_y])
    chart_df = chart_df.reindex(chart_df[resolved_y].abs().sort_values(ascending=False).index)
    chart_df = chart_df.head(20)
    if chart_df.empty:
        return None

    sheet_name = _safe_sheet_name("Gráfico", used_names)
    ws = wb.create_sheet(sheet_name)
    ws.cell(1, 1, resolved_x)
    ws.cell(1, 2, resolved_y)
    for i, (_, row) in enumerate(chart_df.iterrows(), 2):
        ws.cell(i, 1, str(row[resolved_x]))
        ws.cell(i, 2, float(row[resolved_y]) if row[resolved_y] is not None else None)

    last_row = len(chart_df) + 1
    chart_cls = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type, BarChart)
    chart = chart_cls()
    chart.title = f"{resolved_y} por {resolved_x}"
    chart.height = 12
    chart.width = 24
    data_ref = Reference(ws, min_col=2, min_row=1, max_row=last_row)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "D2")
    ws.column_dimensions["A"].width = 30
    return sheet_name


def _build_workbook(
    rows: list[dict[str, Any]],
    title: str,
    var_pct_cols: list[str],
    summary_text: str | None,
    split_by: str | None,
    include_chart: bool,
    chart_type: str,
    chart_x: str | None,
    chart_y: str | None,
    delta_cols: list[str] | None = None,
    highlight_col: str | None = None,
) -> tuple[Any, str | None]:
    """Arma el workbook completo: Resumen (opcional) + Gráfico (opcional) +
    hoja(s) de datos (una sola, o una por valor de split_by). Retorna
    (workbook, error) — error no-None si split_by no es una columna válida."""
    import pandas as pd
    from openpyxl import Workbook

    df = pd.DataFrame(rows)
    for col in df.columns:
        try:
            df[col] = pd.to_numeric(df[col])
        except (ValueError, TypeError):
            pass

    if split_by is not None and split_by not in df.columns:
        return None, (
            f"'split_by'={split_by!r} no es una columna del resultado. "
            f"Columnas disponibles: {list(df.columns)}"
        )

    wb = Workbook()
    wb.remove(wb.active)  # se crean todas las hojas explícitamente abajo
    used_names: set[str] = set()

    if summary_text:
        ws_summary = wb.create_sheet(_safe_sheet_name("Resumen", used_names))
        _write_summary_sheet(ws_summary, summary_text, title)

    if include_chart:
        _add_chart_sheet(wb, df, chart_type, chart_x, chart_y, used_names)

    if split_by:
        for value in sorted(df[split_by].dropna().unique(), key=str):
            subset = df[df[split_by] == value].drop(columns=[split_by])
            ws = wb.create_sheet(_safe_sheet_name(str(value), used_names))
            _write_data_sheet(ws, subset, var_pct_cols, delta_cols, highlight_col)
        na_subset = df[df[split_by].isna()].drop(columns=[split_by])
        if not na_subset.empty:
            ws = wb.create_sheet(_safe_sheet_name("(sin valor)", used_names))
            _write_data_sheet(ws, na_subset, var_pct_cols, delta_cols, highlight_col)
    else:
        ws = wb.create_sheet(_safe_sheet_name(title, used_names))
        _write_data_sheet(ws, df, var_pct_cols, delta_cols, highlight_col)

    wb.active = 0
    return wb, None


class ExportToExcelRequest(BaseModel):
    dataset_id: int = Field(..., description="ID del dataset de Superset")
    report_id: str | None = Field(
        None,
        description=(
            "ID devuelto por irex.compare_periods para exportar exactamente ese "
            "análisis. Si se usa, se ignoran los parámetros de consulta/"
            "comparación y se fuerza modo comparación."
        ),
    )
    title: str = Field("Exportación MCP", description="Título del archivo y hoja de Excel")
    row_limit: int = Field(
        10000, ge=1, le=50000,
        description="Máximo de filas a exportar (default 10.000, máximo 50.000).",
    )

    # ── Modo consulta simple (como irex.query_dataset) ──────────────────────
    metrics: list[str] = Field(
        default_factory=list,
        description=(
            "Métricas a agregar. Si se define sin period_a_filters, usa modo "
            "consulta simple (equivalente a irex.query_dataset). Mismo formato "
            "que irex.query_dataset, incluyendo el alias ' AS <nombre>' para "
            "fórmulas/ratios (ej. 'SUM(plan)/NULLIF(SUM(proyeccion),0) - 1 AS "
            "% Decrecimiento') — usarlo SIEMPRE en métricas calculadas acá, "
            "porque sin alias el header de columna en el Excel queda con la "
            "fórmula SQL cruda en vez de un nombre legible."
        ),
    )
    groupby: list[str] = Field(default_factory=list, description="Columnas de agrupación.")
    filters: list[DatasetFilter] = Field(default_factory=list, description="Filtros WHERE.")
    orderby: list[str] = Field(default_factory=list, description="Ordenamiento.")
    post_filter_expr: str | None = Field(
        None, description="Filtro pandas post-query (mismo formato que irex.query_dataset).",
    )

    # ── Modo comparación de períodos (como irex.compare_periods) ────────────
    base_filters: list[DatasetFilter] = Field(
        default_factory=list,
        description="Filtros comunes a ambos períodos. Usar con period_a/b_filters.",
    )
    period_a_filters: list[DatasetFilter] = Field(
        default_factory=list,
        description=(
            "Filtros del período A (referencia). Si se define, activa modo comparación "
            "(equivalente a irex.compare_periods). Requiere también 'metrics'."
        ),
    )
    period_b_filters: list[DatasetFilter] = Field(
        default_factory=list,
        description="Filtros del período B (actual).",
    )
    period_a_label: str = Field(
        "Período A",
        description="Etiqueta legible del período A para los encabezados de "
        "columna del Excel (ej. 'julio 2026'). Si se pasa 'report_id' y ese "
        "reporte ya trae su propia etiqueta (definida en irex.compare_periods "
        "vía 'period_a_label', o inferida de sus filtros), esa tiene prioridad "
        "— este valor solo se usa como fallback si el reporte no la tiene.",
    )
    period_b_label: str = Field(
        "Período B",
        description="Etiqueta legible del período B (ej. 'agosto 2026'). "
        "Mismo criterio de prioridad que 'period_a_label'.",
    )
    min_variation_pct: float | None = Field(
        None, ge=0,
        description="Solo exportar filas con variación absoluta >= este % (solo modo comparación).",
    )

    # ── Modo SQL (como irex.query_dataset_sql) ──────────────────────────────
    sql: str | None = Field(
        None,
        description=(
            "Consulta SQL (dialecto DuckDB) sobre la tabla 'data' cargada desde Superset. "
            "Si se define, activa modo SQL (equivalente a irex.query_dataset_sql) y tiene "
            "PRIORIDAD sobre los otros modos — ideal para exportar resultados de joins, "
            "pivots, funciones de ventana, etc. que no se pueden expresar con los otros "
            "modos. Usar junto con 'metrics'/'groupby' igual que en query_dataset_sql "
            "(vacíos ambos = trae todas las columnas crudas del dataset). "
            "IMPORTANTE: en este modo, 'row_limit' solo limita las filas del RESULTADO "
            "del sql — usar 'fetch_row_limit' (no 'row_limit') para controlar cuántas "
            "filas crudas se traen de Superset antes de aplicar el sql."
        ),
    )
    fetch_row_limit: int = Field(
        20000, ge=1, le=50000,
        description=(
            "SOLO modo SQL: máximo de filas crudas a traer de Superset hacia la tabla "
            "'data' ANTES de aplicar el sql. Debe ser generoso (default 20.000) cuando el "
            "sql hace JOINs/agregaciones propias — un valor bajo aquí puede hacer que el "
            "sql no encuentre coincidencias aunque la lógica sea correcta, porque faltan "
            "filas fuente. No confundir con 'row_limit', que limita el resultado final."
        ),
    )
    extra_tables: list[ExtraTable] | None = Field(
        None,
        description=(
            "SOLO modo SQL — fuentes ADICIONALES con el MISMO contrato que en "
            "irex.query_dataset_sql: cada una se consulta a Superset por separado "
            "(con sus propios metrics/groupby/filters/jinja_filters/fetch_row_limit, "
            "respetando RLS) y se registra como tabla con su 'name' (ej. 'data2') "
            "junto a 'data', para que el 'sql' haga JOINs entre granularidades o "
            "datasets distintos. Un análisis validado en query_dataset_sql con "
            "extra_tables se exporta pasando el MISMO request más los campos de "
            "presentación. Error si se define sin 'sql'."
        ),
    )

    # ── Presentación del archivo (aplica a los tres modos) ──────────────────
    summary_text: str | None = Field(
        None,
        description=(
            "Análisis en texto (el mismo que se le mostraría al usuario en el chat) para "
            "incluir como primera hoja 'Resumen' del Excel. NUNCA mandar un único párrafo "
            "largo con todos los datos encadenados por comas/paréntesis — se ve como un "
            "bloque de texto denso e ilegible. Estructurar SIEMPRE usando el subconjunto de "
            "markdown soportado: '#'/'##'/'###' para encabezados de sección en negrita; "
            "'-'/'*' o '1.'/'2.' para listas, UNA línea por hallazgo o comparación puntual "
            "en vez de todos los datos en la misma oración; '**texto**' para negrita real "
            "sobre cifras o nombres clave; y párrafos cortos (2-4 oraciones), separados por "
            "una línea en blanco entre secciones. Estructura recomendada:\n"
            "## Contexto\\n"
            "Una o dos oraciones: período, alcance, filtros aplicados.\\n\\n"
            "## Hallazgos principales\\n"
            "- **Producto/categoría X**: +18.3% vs referencia\\n"
            "- **Producto/categoría Y**: -20.0% vs referencia\\n\\n"
            "## Conclusión\\n"
            "Una o dos oraciones de cierre.\n"
            "Usar SIEMPRE que ya se haya redactado un análisis para la respuesta del chat "
            "— así el archivo queda autocontenido sin depender del historial de la "
            "conversación, con el mismo nivel de detalle punto por punto que se le mostró "
            "al usuario (no una versión condensada en un solo párrafo)."
        ),
    )
    split_by: str | None = Field(
        None,
        description=(
            "Columna por la que dividir el resultado en pestañas separadas — una hoja por "
            "cada valor distinto (ej. 'formato' o 'familia'), en vez de una única hoja larga. "
            "Debe ser una de las columnas presentes en el resultado final (una columna de "
            "'groupby', o una columna seleccionada por el 'sql' en modo SQL). Usar cuando "
            "el usuario pida el archivo 'organizado por' o 'separado por' alguna dimensión."
        ),
    )
    include_chart: bool = Field(
        False,
        description=(
            "Si True, agrega una hoja 'Gráfico' con un gráfico nativo de Excel (interactivo, "
            "no una imagen) construido a partir de las 20 filas de mayor magnitud del "
            "resultado. Usar cuando el usuario pida que el Excel 'incluya un gráfico' o "
            "'se vea más visual'."
        ),
    )
    chart_type: Literal["bar", "line", "pie"] = Field(
        "bar", description="Tipo de gráfico cuando include_chart=True.",
    )
    chart_x: str | None = Field(
        None,
        description=(
            "Columna para el eje X / categorías del gráfico (ej. 'articulo', 'formato'). "
            "Si no se especifica, se usa automáticamente la primera columna no numérica."
        ),
    )
    chart_y: str | None = Field(
        None,
        description=(
            "Columna numérica para el eje Y / valores del gráfico (ej. 'AVG(precio)', "
            "una columna de variación %). Si no se especifica, se usa automáticamente la "
            "primera columna numérica."
        ),
    )
    jinja_filters: list[JinjaFilterOverride] | None = Field(
        None,
        description="AVANZADO — usar solo cuando el dataset tiene columnas "
        "calculadas en su SQL vía filter_values()/get_filters() (típicamente "
        "lógica de 'año seleccionado vs año anterior' u otra dependiente de un "
        "filtro nativo del dashboard) y se necesita fijar ese valor SIN pedirle "
        "al usuario que cambie el filtro en el dashboard. Cada entry simula la "
        "selección de un filtro nativo: {'column': '<target del filtro, ej. "
        "anio_ids>', 'value': [<valor>]}. Aplica a los TRES modos (consulta, "
        "comparación y SQL) — en modo comparación se aplica idéntico a ambos "
        "períodos. FUENTE CONFIRMADA: 'native_filter_targets' en la respuesta de "
        "irex.get_query_context/get_dashboard_dataset_context — usar SOLO un "
        "'column' que aparezca ahí con usable_as='jinja_filters'. Si un dashboard "
        "tiene filtros activos clasificados así y no se pasan acá, el Excel "
        "exportado NO los reflejará (son inaplicables como 'filters'/"
        "'base_filters' normales).",
    )


@tool(
    name="irex.export_to_excel",
    description=(
        "Ejecuta una consulta, comparación de períodos, o análisis SQL y exporta "
        "el resultado completo a un archivo Excel descargable — sin límite de "
        "tokens. Usar cuando el usuario quiere ver o descargar TODOS los registros "
        "de un análisis que es demasiado grande para mostrar en el chat. "
        "Soporta tres modos: "
        "(1) Modo consulta: definir 'metrics' — igual que irex.query_dataset; "
        "(2) Modo comparación: definir 'metrics' + 'period_a_filters' + "
        "'period_b_filters' — igual que irex.compare_periods; "
        "(3) Modo SQL: definir 'sql' — igual que irex.query_dataset_sql, incluidas "
        "las 'extra_tables' para JOINs multi-fuente (exportar un análisis ya "
        "validado ahí = repetir el mismo request y sumar título/summary/split_by), para "
        "exportar joins/pivots/funciones de ventana que los otros modos no cubren. "
        "En modo SQL, 'fetch_row_limit' (no 'row_limit') controla cuántas filas crudas "
        "se traen antes de aplicar el sql — subirlo si el sql hace JOINs/agregaciones "
        "propias y el resultado sale más chico de lo esperado. "
        "Presentación (independiente del modo): 'summary_text' agrega una hoja 'Resumen' "
        "con el análisis en texto; 'split_by' divide el resultado en una pestaña por cada "
        "valor de una columna; 'include_chart' agrega una hoja 'Gráfico' con un gráfico "
        "nativo de Excel interactivo. "
        "Soporta 'jinja_filters' (igual que irex.query_dataset) para datasets con "
        "columnas calculadas vía filter_values()/get_filters() en su SQL virtual — "
        "usar cuando 'get_query_context'/'get_dashboard_dataset_context' clasifiquen "
        "un filtro nativo del dashboard como usable_as='jinja_filters'; sin esto, "
        "el Excel exportado NO reflejará esos filtros aunque estén activos en el "
        "dashboard. "
        "Si se pasa 'report_id' devuelto por irex.compare_periods, exporta "
        "EXACTAMENTE ese análisis guardado (mismo cálculo, mismo orden de "
        "'sort_by', mismas etiquetas de período) sin reconstruir filtros/"
        "métricas manualmente — el Excel de comparación incluye columnas de "
        "valor por período, variación %, cambio absoluto y % de contribución "
        "(para métricas aditivas), con formato condicional verde/rojo en "
        "variación y cambio absoluto, y un borde resaltando la fila de mayor "
        "aumento y mayor disminución. "
        "Retorna una URL de descarga directa (válida 2 horas)."
    ),
    class_permission_name="SQLLab",
    method_permission_name="execute_sql_query",
)
def export_to_excel(request: ExportToExcelRequest) -> dict[str, Any]:
    _cleanup_old_exports()
    _EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    eff_dataset_id = request.dataset_id
    eff_metrics = request.metrics
    eff_groupby = request.groupby
    eff_filters = request.filters
    eff_base_filters = request.base_filters
    eff_period_a_filters = request.period_a_filters
    eff_period_b_filters = request.period_b_filters
    eff_jinja_filters = request.jinja_filters
    eff_min_variation_pct = request.min_variation_pct
    eff_sort_by = "variation_pct"
    eff_period_a_label = request.period_a_label
    eff_period_b_label = request.period_b_label
    use_sql = bool(request.sql)
    use_compare = bool(request.period_a_filters)

    if request.report_id:
        cached = report_cache.get_report(request.report_id)
        if cached is None:
            return {
                "status": "error",
                "error": (
                    "report_id inválido, vencido o de otro usuario. Repetir "
                    "irex.compare_periods y usar el nuevo report_id."
                ),
            }
        eff_dataset_id = cached["dataset_id"]
        eff_metrics = cached.get("metrics", [])
        eff_groupby = cached.get("groupby", [])
        eff_base_filters = [
            DatasetFilter(**f) for f in cached.get("base_filters", [])
        ]
        eff_period_a_filters = [
            DatasetFilter(**f) for f in cached.get("period_a_filters", [])
        ]
        eff_period_b_filters = [
            DatasetFilter(**f) for f in cached.get("period_b_filters", [])
        ]
        cached_jinja = cached.get("jinja_filters")
        eff_jinja_filters = (
            [JinjaFilterOverride(**f) for f in cached_jinja]
            if cached_jinja
            else None
        )
        eff_min_variation_pct = cached.get("min_variation_pct")
        # Reportes guardados ANTES de este cambio no tienen 'sort_by' ni
        # etiquetas de período — caen a los mismos defaults que ya existían
        # (compatibilidad hacia atrás, sin romper reportes viejos).
        eff_sort_by = cached.get("sort_by", "variation_pct")
        eff_period_a_label = cached.get("period_a_label") or request.period_a_label
        eff_period_b_label = cached.get("period_b_label") or request.period_b_label
        use_sql = False
        use_compare = True

    # extra_tables solo tiene sentido en modo SQL — rechazarlo explícito en vez
    # de ignorarlo en silencio (el LLM creería que el JOIN multi-fuente aplicó).
    if request.extra_tables and not use_sql:
        return {
            "status": "error",
            "error": (
                "'extra_tables' solo aplica al modo SQL — definir también 'sql' "
                "(mismo contrato que irex.query_dataset_sql), o quitar "
                "'extra_tables' si el export es de agregación/comparación simple."
            ),
        }

    rows: list[dict[str, Any]] = []
    var_pct_cols: list[str] = []
    delta_cols: list[str] = []
    highlight_col: str | None = None
    # Diagnósticos de completitud del modo SQL — se propagan a la respuesta
    # final (y al Excel) para que un archivo calculado sobre fuentes truncadas
    # nunca se presente como exacto.
    sql_diagnostics: dict[str, Any] = {}

    if use_sql:
        sql_result = execute_sql_analysis(
            dataset_id=eff_dataset_id,
            metrics=eff_metrics,
            groupby=eff_groupby,
            filters=eff_filters,
            fetch_row_limit=request.fetch_row_limit,
            sql=request.sql,
            result_row_limit=request.row_limit,
            jinja_filters=eff_jinja_filters,
            extra_tables=request.extra_tables,
        )
        if sql_result.get("status") == "error":
            return sql_result
        if sql_result.get("status") == "no_data":
            return {
                "status": "no_data",
                "row_count": 0,
                "message": "La consulta no devolvió filas. No se generó el archivo.",
            }
        rows = sql_result.get("rows", [])
        for key in (
            "result_exact",
            "source_row_count",
            "source_truncated",
            "source_truncated_warning",
            "incomplete_reason",
            "extra_sources",
        ):
            if key in sql_result:
                sql_diagnostics[key] = sql_result[key]

    elif use_compare:
        from .dashboard_dataset_context import _reject_hidden_columns

        referenced = (
            set(eff_groupby)
            | {f.column for f in eff_base_filters}
            | {f.column for f in eff_period_a_filters}
            | {f.column for f in eff_period_b_filters}
        )
        hidden_error = _reject_hidden_columns(eff_dataset_id, referenced)
        if hidden_error:
            return {"status": "error", "error": hidden_error}

        if not eff_metrics:
            return {
                "status": "error",
                "error": "Modo comparación requiere 'metrics' además de 'period_a_filters'.",
            }
        metrics_parsed = [_parse_metric(m) for m in eff_metrics]
        metric_labels: list[str] = [
            m if isinstance(m, str) else m.get("label", str(m))
            for m in metrics_parsed
        ]
        additive_labels = {
            label
            for raw, label in zip(eff_metrics, metric_labels)
            if _is_additive_metric(raw)
        }

        # Mismo cálculo que irex.compare_periods (_build_comparison) — antes
        # este modo reimplementaba el cruce A/B por su cuenta, con menos
        # columnas ('_delta' faltante) y sin respetar el 'sort_by' original,
        # lo que hacía que el Excel divergiera de la tabla mostrada en el
        # chat. A diferencia de compare_periods (donde row_limit acota el
        # resultado YA cruzado), acá se quiere el detalle completo hasta
        # row_limit filas por período — se pasa como fetch_limit explícito.
        comparison = _build_comparison(
            eff_dataset_id, metrics_parsed, metric_labels, additive_labels,
            eff_groupby, eff_base_filters, eff_period_a_filters, eff_period_b_filters,
            eff_sort_by, eff_min_variation_pct, eff_jinja_filters,
            fetch_limit=request.row_limit,
        )
        if comparison["error"]:
            return {"status": "error", "error": comparison["error"]}

        rows, var_pct_cols, delta_cols = _format_comparison_rows(
            comparison["combined"],
            metric_labels,
            eff_period_a_label or "Período A",
            eff_period_b_label or "Período B",
        )
        # Resaltar mayor aumento/mayor disminución sobre la PRIMERA métrica
        # (criterio simple y genérico — con varias métricas, el resto sigue
        # teniendo su propio formato condicional por celda, solo no define
        # qué fila se resalta con el borde).
        highlight_col = delta_cols[0] if delta_cols else None

    else:
        from .dashboard_dataset_context import _reject_hidden_columns

        referenced = set(eff_groupby) | {f.column for f in eff_filters}
        hidden_error = _reject_hidden_columns(eff_dataset_id, referenced)
        if hidden_error:
            return {"status": "error", "error": hidden_error}

        if not eff_metrics:
            return {
                "status": "error",
                "error": (
                    "Se requiere 'metrics' para modo consulta, o 'metrics' + "
                    "'period_a_filters' para modo comparación."
                ),
            }
        from superset.commands.chart.data.get_data_command import ChartDataCommand
        from superset.common.query_context_factory import QueryContextFactory

        metrics_parsed = [_parse_metric(m) for m in eff_metrics]
        adhoc_filters = [_parse_filter(f) for f in eff_filters]
        orderby = (
            _parse_orderby(request.orderby, eff_metrics) if request.orderby else []
        )

        with _jinja_filters_context(eff_jinja_filters) as forced:
            factory = QueryContextFactory()
            qc = factory.create(
                datasource={"id": eff_dataset_id, "type": "table"},
                force=forced,
                queries=[{
                    "filters": adhoc_filters,
                    "columns": eff_groupby,
                    "metrics": metrics_parsed,
                    "row_limit": request.row_limit,
                    "orderby": orderby,
                }],
                form_data={"datasource": f"{eff_dataset_id}__table", "viz_type": "table"},
            )
            cmd = ChartDataCommand(qc)
            cmd.validate()
            try:
                result = cmd.run()
            except Exception as exc:
                return {"status": "error", "error": str(exc)}

        qr = (result or {}).get("queries", [{}])[0]
        if qr.get("error"):
            return {"status": "error", "error": qr["error"]}
        rejected = qr.get("rejected_filters", [])
        if rejected:
            return {
                "status": "error",
                "error": "Columna(s) no encontrada(s): "
                + ", ".join(str(r.get("column", r)) for r in rejected),
            }
        rows = _sanitize_rows(qr.get("data", []))

        if request.post_filter_expr and rows:
            rows, filter_err = _apply_post_filter(rows, request.post_filter_expr)
            if filter_err:
                return {"status": "error", "error": filter_err}

    if not rows:
        return {
            "status": "no_data",
            "row_count": 0,
            "message": "La consulta no devolvió filas. No se generó el archivo.",
        }

    # Si alguna fuente quedó truncada, el caveat viaja DENTRO del archivo
    # (hoja Resumen) — el Excel circula sin el contexto del chat, así que la
    # respuesta del tool no alcanza como único aviso.
    effective_summary = request.summary_text
    truncation_warning = sql_diagnostics.get("source_truncated_warning")
    if truncation_warning:
        caveat = (
            "## ⚠️ Datos posiblemente incompletos\n"
            f"{truncation_warning}\n\n"
        )
        effective_summary = caveat + (effective_summary or "")

    try:
        wb, build_error = _build_workbook(
            rows=rows,
            title=request.title,
            var_pct_cols=var_pct_cols,
            summary_text=effective_summary,
            split_by=request.split_by,
            include_chart=request.include_chart,
            chart_type=request.chart_type,
            chart_x=request.chart_x,
            chart_y=request.chart_y,
            delta_cols=delta_cols,
            highlight_col=highlight_col,
        )
    except Exception as exc:
        return {"status": "error", "error": f"Error generando Excel: {exc}"}
    if build_error:
        return {"status": "error", "error": build_error}

    token = secrets.token_urlsafe(16)
    output_path = _EXPORT_DIR / f"{token}.xlsx"
    wb.save(output_path)

    download_url = f"{_get_export_base_url()}/api/mcp/download/{token}"

    extras = []
    if request.summary_text:
        extras.append("hoja de resumen")
    if request.include_chart:
        extras.append("gráfico")
    if request.split_by:
        extras.append(f"dividido por '{request.split_by}'")
    extras_msg = f" ({', '.join(extras)})" if extras else ""

    response: dict[str, Any] = {
        "status": "success",
        "row_count": len(rows),
        "download_url": download_url,
        "expires_in": "2 horas",
        "message": (
            f"Excel generado con {len(rows):,} filas{extras_msg}. "
            f"Descarga disponible 2 horas: {download_url}"
        ),
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }
    response.update(sql_diagnostics)
    if truncation_warning:
        response["message"] += (
            " ⚠️ ATENCIÓN: una o más fuentes quedaron truncadas — el archivo "
            "incluye el aviso en su hoja Resumen y NO debe presentarse como "
            "exacto (ver source_truncated_warning)."
        )
    return response
